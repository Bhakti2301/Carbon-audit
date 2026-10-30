
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = (
    BASE_DIR
    / "data"
    / "source"
    / "ghg-conversion-factors-2026-flat-format-revised.xlsx"
)

SHEET_NAME = "Factors by Category"
HEADER_ROW = 6

SEARCH_GROUPS = {
    "Coal and coke": ["coal", "coke"],
    "Diesel and gas oil": ["diesel", "gas oil"],
    "Natural gas": ["natural gas"],
    "Electricity": ["electricity", "electricity: "],
    "Petrol": ["petrol"],
    "Steel and metals": ["steel", "iron", "metal"],
    "Freight and transport": [
        "freight", "road", "vehicle", "lorry", "truck"
    ],
    "Refrigerants and SF6": [
        "refrigerant", "sf6", "sulphur hexafluoride",
        "sulfur hexafluoride"
    ],
    "Livestock and agriculture": [
        "cattle", "dairy cow", "livestock", "fertiliser",
        "fertilizer", "nitrogen", "agriculture"
    ],
    "Waste and water": ["waste", "water"],
    "Flights and business travel": [
        "flight", "air travel", "business travel"
    ],
    "Purchased goods and materials": [
        "purchased goods", "material", "steel"
    ],
}


def clean(value):
    if value is None:
        return ""

    return str(value).strip()


def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Workbook not found: {EXCEL_FILE}"
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True
    )

    try:
        worksheet = workbook[SHEET_NAME]

        headers = [
            clean(cell)
            for cell in next(
                worksheet.iter_rows(
                    min_row=HEADER_ROW,
                    max_row=HEADER_ROW,
                    values_only=True
                )
            )
        ]

        records = []

        for values in worksheet.iter_rows(
            min_row=HEADER_ROW + 1,
            values_only=True
        ):
            if not any(value is not None for value in values):
                continue

            padded = list(values[:len(headers)])
            padded.extend([None] * (len(headers) - len(padded)))

            record = dict(zip(headers, padded))

            records.append(record)

        print("=" * 80)
        print("DEFRA 2026 CATEGORY SUMMARY")
        print("=" * 80)

        print(f"Data rows inspected: {len(records)}")

        for column in [
            "Scope",
            "Level 1",
            "Level 2",
            "GHG/Unit",
            "UOM"
        ]:
            counter = Counter(
                clean(row.get(column))
                for row in records
                if clean(row.get(column))
            )

            print(f"\n--- {column}: most common values ---")

            for value, count in counter.most_common(25):
                print(f"{value}: {count}")

        print("\n" + "=" * 80)
        print("RELEVANT ACTIVITY FACTORS")
        print("=" * 80)

        for group_name, keywords in SEARCH_GROUPS.items():
            matches = []

            for row in records:
                searchable = " ".join(
                    clean(row.get(column))
                    for column in [
                        "Level 1",
                        "Level 2",
                        "Level 3",
                        "Level 4",
                        "Column Text",
                        "UOM",
                        "GHG/Unit"
                    ]
                ).casefold()

                if any(keyword.casefold() in searchable for keyword in keywords):
                    matches.append(row)

            # Prefer complete aggregate CO2e factors for inspection.
            # Component-specific factors remain visible but are not
            # automatically added to the aggregate factor.
            matches.sort(
                key=lambda row: (
                    0
                    if clean(row.get("GHG/Unit")).casefold()
                    in {"kg co2e", "kgco2e"}
                    else 1,
                    clean(row.get("Level 1")),
                    clean(row.get("Level 2")),
                    clean(row.get("Level 3")),
                    clean(row.get("UOM")),
                    clean(row.get("ID"))
                )
            )

            print(f"\n--- {group_name}: {len(matches)} matches ---")

            # Limit each group so terminal output stays manageable.
            for row in matches[:8]:
                print(
                    f"ID={clean(row.get('ID'))} | "
                    f"Scope={clean(row.get('Scope'))} | "
                    f"Category="
                    f"{clean(row.get('Level 1'))} > "
                    f"{clean(row.get('Level 2'))} > "
                    f"{clean(row.get('Level 3'))} > "
                    f"{clean(row.get('Level 4'))} | "
                    f"Description={clean(row.get('Column Text'))} | "
                    f"Activity unit={clean(row.get('UOM'))} | "
                    f"Output={clean(row.get('GHG/Unit'))} | "
                    f"Factor={clean(row.get('GHG Conversion Factor 2026'))}"
                )

            if len(matches) > 8:
                print(f"... {len(matches) - 8} additional matches omitted")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()
