
from pathlib import Path
import csv
import re
import sys
from collections import Counter

from openpyxl import load_workbook


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

SOURCE_FILE = (
    BASE_DIR
    / "data"
    / "source"
    / "ghg-conversion-factors-2026-flat-format-revised.xlsx"
)

OUTPUT_FILE = (
    BASE_DIR
    / "data"
    / "source"
    / "defra_2026_factors_normalized.csv"
)

SHEET_NAME = "Factors by Category"
HEADER_ROW = 6


# ============================================================
# SOURCE METADATA
# ============================================================

SOURCE_ID = "defra_desnz"
SOURCE_NAME = "UK Government GHG Conversion Factors"
SOURCE_URL = (
    "https://www.gov.uk/government/publications/"
    "greenhouse-gas-reporting-conversion-factors-2026"
)


# ============================================================
# HELPERS
# ============================================================

def clean(value):
    """Convert a cell to a trimmed string, preserving blank cells."""

    if value is None:
        return ""

    return str(value).strip()


def normalize_scope(scope_label):
    """Extract the scope number without changing the original label."""

    match = re.search(
        r"scope\s*([123])",
        clean(scope_label),
        re.IGNORECASE
    )

    return match.group(1) if match else ""


def normalize_factor_value(value):
    """
    Return the factor as text suitable for CSV.

    Blank factors stay blank. They are never changed to zero.
    """

    if value is None:
        return "", "missing"

    if isinstance(value, str) and not value.strip():
        return "", "missing"

    try:
        number = float(value)
    except (TypeError, ValueError):
        return clean(value), "non_numeric"

    # Preserve zero as published, but flag it for later review.
    if number == 0:
        return str(value), "zero_requires_review"

    return str(value), "present"


# ============================================================
# IMPORTER
# ============================================================

def main():
    if not SOURCE_FILE.is_file():
        print("ERROR: The source workbook was not found.")
        print(f"Expected path: {SOURCE_FILE}")
        sys.exit(1)

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    print("=" * 72)
    print("DEFRA FLAT-FILE NORMALIZER")
    print("=" * 72)
    print(f"Source: {SOURCE_FILE}")
    print(f"Output: {OUTPUT_FILE}")

    workbook = load_workbook(
        filename=SOURCE_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Required worksheet '{SHEET_NAME}' was not found. "
                f"Available worksheets: {workbook.sheetnames}"
            )

        # Read metadata from the workbook rather than assuming the
        # current reporting year or source release.
        front_page = workbook["Front page"]

        dataset_year = clean(
            front_page["E5"].value
            or front_page["C18"].value
        )

        source_release = clean(
            front_page["C5"].value
            or front_page["C13"].value
        )

        updated_value = front_page["C14"].value
        source_updated_at = clean(updated_value)

        if not dataset_year.isdigit() or len(dataset_year) != 4:
            raise ValueError(
                "Could not identify a four-digit dataset year "
                "from the workbook metadata."
            )

        worksheet = workbook[SHEET_NAME]

        header_values = next(
            worksheet.iter_rows(
                min_row=HEADER_ROW,
                max_row=HEADER_ROW,
                values_only=True
            )
        )

        headers = [clean(value) for value in header_values]

        required_headers = [
            "ID",
            "Scope",
            "Level 1",
            "Level 2",
            "Level 3",
            "Level 4",
            "Column Text",
            "UOM",
            "GHG/Unit",
            f"GHG Conversion Factor {dataset_year}"
        ]

        missing_headers = [
            name for name in required_headers
            if name not in headers
        ]

        if missing_headers:
            raise ValueError(
                "The workbook structure doesn't match the expected "
                f"format. Missing headers: {missing_headers}"
            )

        output_fields = [
            "source_id",
            "source_name",
            "source_url",
            "source_release",
            "source_updated_at",
            "dataset_year",
            "factor_id",
            "scope_label",
            "scope_number",
            "category_level_1",
            "category_level_2",
            "category_level_3",
            "category_level_4",
            "column_text",
            "activity_unit",
            "ghg_unit",
            "factor_value",
            "factor_status"
        ]

        factor_column = f"GHG Conversion Factor {dataset_year}"

        row_count = 0
        scope_counts = Counter()
        unit_counts = Counter()
        status_counts = Counter()

        with open(
            OUTPUT_FILE,
            "w",
            newline="",
            encoding="utf-8-sig"
        ) as output_file:

            writer = csv.DictWriter(
                output_file,
                fieldnames=output_fields
            )

            writer.writeheader()

            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=HEADER_ROW + 1,
                    values_only=True
                ),
                start=HEADER_ROW + 1
            ):
                source_row = dict(zip(headers, values))

                factor_id = clean(source_row.get("ID"))

                # Ignore title, notes and empty rows.
                if not factor_id:
                    continue

                factor_value, factor_status = normalize_factor_value(
                    source_row.get(factor_column)
                )

                scope_label = clean(source_row.get("Scope"))

                normalized = {
                    "source_id": SOURCE_ID,
                    "source_name": SOURCE_NAME,
                    "source_url": SOURCE_URL,
                    "source_release": source_release,
                    "source_updated_at": source_updated_at,
                    "dataset_year": dataset_year,

                    # Preserve the original factor identifier.
                    "factor_id": factor_id,

                    # Keep both source label and parsed scope.
                    "scope_label": scope_label,
                    "scope_number": normalize_scope(scope_label),

                    # Preserve DEFRA's category structure.
                    "category_level_1": clean(
                        source_row.get("Level 1")
                    ),
                    "category_level_2": clean(
                        source_row.get("Level 2")
                    ),
                    "category_level_3": clean(
                        source_row.get("Level 3")
                    ),
                    "category_level_4": clean(
                        source_row.get("Level 4")
                    ),

                    "column_text": clean(
                        source_row.get("Column Text")
                    ),
                    "activity_unit": clean(
                        source_row.get("UOM")
                    ),
                    "ghg_unit": clean(
                        source_row.get("GHG/Unit")
                    ),

                    # Never replace a blank factor with zero.
                    "factor_value": factor_value,
                    "factor_status": factor_status
                }

                writer.writerow(normalized)

                row_count += 1
                scope_counts[scope_label or "(blank)"] += 1
                unit_counts[
                    normalized["activity_unit"] or "(blank)"
                ] += 1
                status_counts[factor_status] += 1

        print("\nIMPORT COMPLETE")
        print(f"Dataset year: {dataset_year}")
        print(f"Source release: {source_release}")
        print(f"Source updated: {source_updated_at}")
        print(f"Records exported: {row_count}")
        print(f"Output size: {OUTPUT_FILE.stat().st_size:,} bytes")

        print("\nFactor status counts:")
        for status, count in sorted(status_counts.items()):
            print(f"  {status}: {count}")

        print("\nScope counts:")
        for scope, count in sorted(scope_counts.items()):
            print(f"  {scope}: {count}")

        print("\nMost common activity units:")
        for unit, count in unit_counts.most_common(15):
            print(f"  {unit}: {count}")

        print("\nIMPORTANT:")
        print(
            "This is a normalized copy of the source data, "
            "not yet the application's calculation-factor file."
        )
        print(
            "No industry/process mappings were invented. "
            "Missing values remain blank."
        )

    finally:
        workbook.close()


if __name__ == "__main__":
    main()
