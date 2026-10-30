
from pathlib import Path
import sys

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("The openpyxl package is missing.")
    print("Run: python -m pip install openpyxl")
    sys.exit(1)


# Locate the downloaded official workbook.
BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = (
    BASE_DIR
    / "data"
    / "source"
    / "ghg-conversion-factors-2026-flat-format-revised.xlsx"
)


def display_value(value, limit=140):
    """Make worksheet values readable in terminal output."""

    text = str(value).replace("\n", " ").replace("\r", " ")

    if len(text) > limit:
        return text[:limit] + "..."

    return text


def main():
    if not EXCEL_FILE.is_file():
        print("ERROR: DEFRA Excel file was not found.")
        print(f"Expected location:\n{EXCEL_FILE}")
        print("\nMove the downloaded workbook to that location.")
        sys.exit(1)

    print("=" * 80)
    print("DEFRA DATASET INSPECTION")
    print("=" * 80)
    print(f"File: {EXCEL_FILE.name}")
    print(f"Size: {EXCEL_FILE.stat().st_size:,} bytes")

    try:
        workbook = load_workbook(
            filename=EXCEL_FILE,
            read_only=True,
            data_only=True
        )
    except Exception as error:
        print(f"ERROR: Could not open workbook: {error}")
        sys.exit(1)

    try:
        print("\nWORKSHEETS")

        for worksheet in workbook.worksheets:
            print(
                f"- {worksheet.title}: "
                f"{worksheet.max_row} rows x "
                f"{worksheet.max_column} columns"
            )

        # Show sample rows from each sheet to discover
        # its real structure before creating an importer.
        for worksheet in workbook.worksheets:
            print("\n" + "=" * 80)
            print(f"SHEET: {worksheet.title}")
            print("=" * 80)

            nonempty_rows = 0

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1
            ):
                populated_cells = []

                for column_number, value in enumerate(
                    row,
                    start=1
                ):
                    if value is None:
                        continue

                    if isinstance(value, str) and not value.strip():
                        continue

                    column_letter = get_column_letter(
                        column_number
                    )

                    populated_cells.append(
                        f"{column_letter}="
                        f"{display_value(value)}"
                    )

                if not populated_cells:
                    continue

                print(f"\nRow {row_number}:")

                # Prevent exceptionally wide rows from flooding
                # the terminal with output.
                for start in range(
                    0,
                    len(populated_cells),
                    8
                ):
                    group = populated_cells[start:start + 8]
                    print("  " + " | ".join(group))

                nonempty_rows += 1

                # Enough initial rows to locate headers and
                # understand the dataset's layout.
                if nonempty_rows >= 12:
                    break

            if nonempty_rows == 0:
                print("No non-empty rows found in this worksheet.")

        print("\n" + "=" * 80)
        print("INSPECTION FINISHED")
        print("=" * 80)
        print(
            "No source data has been modified. "
            "This script only reads the workbook."
        )

    finally:
        workbook.close()


if __name__ == "__main__":
    main()
